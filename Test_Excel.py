
from M_Config import *

import os
from datetime import datetime

import openpyxl

#my_path = "C:\Users\TP\Desktop\Book1.xlsx"
#my_wb_obj = openpyxl.load_workbook(my_path)

my_wb = openpyxl.Workbook()

#my_wb.create_sheet(index = 1 , title = "new sheet")


my_sheet = my_wb.active

#print(my_sheet_obj.max_row)
#my_cell_obj = my_sheet_obj.cell(row = 1, column = 1)
#print(my_cell_obj.value)
#print(sheet_obj.max_column)

c1 = my_sheet.cell(row = 1, column = 1)
c1.value = "Aadrika"
c2 = my_sheet.cell(row= 1 , column = 2)
c2.value = "Adwaita"
c3 = my_sheet['A2']
c3.value = "Satyajit"
# B2 = column = 2 & row = 2.
c4 = my_sheet['B2']
c4.value = "Bivas"


date_time = datetime.now().strftime("%Y %m %d, %H %M")
path = os.path.join( GetPathToDataFolder(), "Statistics " + date_time + ".xlsx" )
my_wb.save( path )